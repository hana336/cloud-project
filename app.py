from flask import Flask, render_template, send_from_directory, request, redirect, url_for, flash
from PyPDF2 import PdfReader
from scrape import scrape_files
from docx import Document
from docx.shared import RGBColor
from docx.enum.text import WD_COLOR_INDEX
from werkzeug.utils import secure_filename
import os
import glob
import sys
import re
import fitz
import requests
from google.cloud import storage



app = Flask(__name__)
app.config['UPLOAD_FOLDER'] = os.path.join(os.getcwd(), 'data/raw')
app.config['ALLOWED_EXTENSIONS'] = {'pdf', 'docx'}
app.config['MAX_CONTENT_LENGTH'] = 16 * 1024 * 1024 #16MB
app.secret_key = 'secret'

DOWNLOAD_FOLDER = app.config['UPLOAD_FOLDER']
os.makedirs(DOWNLOAD_FOLDER, exist_ok=True)

CLASSIFICATION_TREE = {
        "Education":["education", "learning", "students", "teacher"],
        "Technology":["technology", "AI", "cloud", "data"],
        "Medical":["health", "medicine", "doctor", "treatment"],
        "finance":["money","bank","finance","investment"]
}

def allowed_file(filename):
    return '.' in filename and \
        filename.rsplit('.', 1)[1].lower() in app.config['ALLOWED_EXTENSIONS']

# وظيفة التصنيف 
def classify_file(filepath):
    if not os.path.exists(filepath):
        return "Uncategorized"
    content = ""
    try:
        if filepath.lower().endswith('.pdf'):
            reader = PdfReader(filepath)
            content = " ".join(page.extract_text() or "" for page in reader.pages)
        elif filepath.lower().endswith('.docx'):
            doc = docx.Document(filepath)
            content = " ".join(para.text for para in doc.paragraphs)
        else:
            with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
                content = f.read()
        
        content = content.lower()
        for category, keywords in CLASSIFICATION_TREE.items():
            if any(keyword in content for keyword in keywords):
                return category
    except Exception as e:
        print(f"Classification error for {filepath}: {e}")
    return "Uncategorized"



def download_file(url, filename):
    response = requests.get(url)
    if response.status_code == 200:
        path = os.path.join(DOWNLOAD_FOLDER, filename)
        with open(path, "wb") as f:
            f.write(response.content)
        print(f"{filename} downloaded successfully.")
        return True
    else:
        print(f"Failed to download. Status code:{response.status_code}")
        return False  



def extract_title_from_metadata(file_path):
    title = "Unknown"
    try:
        if file_path.lower().endswith(".pdf"):
            reader = PdfReader(file_path)
            info = reader.metadata
            if info and info.title:
                title = info.title
        elif file_path.lower().endswith(".docx"):
            doc = Document(file_path)
            props = doc.core_properties
            if props.title:
                title = props.title
    except Exception as e:
        print(f"Error extracting title: {e}")
    return title                        


# احصائيات الملفات
def get_file_stats():
    files = os.listdir(DOWNLOAD_FOLDER)
    stats = {'total_files': len(files),
             'total_size':f"{sum(os.path.getsize(os.path.join(DOWNLOAD_FOLDER, f)) for f in files) / (1024*1024): .2f} MB",
             'categories': {}}
    for file in files:
        category = classify_file(os.path.join(DOWNLOAD_FOLDER, file))
        stats['categories'][category] = stats['categories'].get(category,0) + 1
    return stats


# وظيفة البحث في الملفات من نوع pdf مع تمييز النص 
def highlight_text_in_pdf(pdf_path, keyword, output_folder="static/highlighted"):
    os.makedirs(output_folder, exist_ok=True)
    output_path = os.path.join(output_folder, os.path.basename(pdf_path))
    doc = fitz.open(pdf_path)
    for page in doc:
        text_instances = page.search_for(keyword)
        for inst in text_instances:
            page.add_highlight_annot(inst)
    doc.save(output_path)
    return output_path


# وظيفة البحث في الملفات من نوع word مع تمييز النص 
def highlight_text_in_docx(docx_path, keyword, output_folder="static/highlighted"):
    os.makedirs(output_folder, exist_ok=True)
    output_path = os.path.join(output_folder, os.path.basename(docx_path))
    doc = Document(docx_path)
    for para in doc.paragraphs:
        if keyword in para.text:
            for run in para.runs:
                if keyword in run.text:
                    run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    doc.save(output_path)
    return output_path


#رفع الملف الى google cloud storage
def upload_file_to_bucket(file_path, bucket_name, destination_name):
    storage_client = storage.Client()
    bucket = storage_client.bucket(bucket_name)
    blob = bucket.blob(destination_name)
    blob.upload_from_filename(file_path)
    print(f"تم رفع {file_path} الى {bucket_name}/ {destination_name}")
    


@app.route('/', methods=["GET", "POST"])
def home():
    sort_order = request.args.get("sort", "asc")
    query = request.args.get("q","")
    files = os.listdir(DOWNLOAD_FOLDER)
    files.sort(reverse=sort_order == 'desc')

    file_data = []
    for f in files:
        path = os.path.join(DOWNLOAD_FOLDER, f)
        category = classify_file(path)
        title = extract_title_from_metadata(path)
        file_data.append({'name': f, 'category': category, 'title':title})


    if query:
        return redirect(url_for('search', q=query, sort=sort_order))
        
    if request.method == "POST":
        url = request.form.get("url")
        filename = request.form.get("filename")
        success = download_file(url, filename)
        if success:
            flash("File downloaded successfully.", "success")
        else:
            flash("Failed to download file.", "danger")
        return redirect(url_for("home", sort=sort_order,q=query))
    
    return render_template('index.html', files=file_data, sort_order= sort_order, query=query)

@app.route('/stats')
def stats():
    try:
        files = os.listdir(DOWNLOAD_FOLDER)
        if not files:
            flash("No files found to analyze", "info")
            return redirect(url_for('home'))
        total_size = 0
        for f in files:
            try: total_size += os.path.getsize(os.path.join(DOWNLOAD_FOLDER, f))
            except Exception as e:
                print(f"Error getting size for {f}:{str(e)}")
                continue
        stats = {'total_files': len(files),
             'total_size':f"{total_size/(1024*1024): .2f} MB",
             'categories': {}}
        for file in files:
            category = classify_file(os.path.join(DOWNLOAD_FOLDER, file))
            stats['categories'][category] = stats['categories'].get(category, 0) + 1
        return render_template('stats.html', stats=stats)    
    except Exception as e:
        flash(f"Error generating statistics: {str(e)}", "danger")
    return redirect(url_for('home'))


@app.route('/scrape')
def scrape_route():
    try:
        test_url = "https://example.com/sample.pdf"
        download_file(test_url, "test_scrape.pdf")
        flash(f"تم تنزيل ملف تجريبي بنجاح", "success")
    except Exception as e:
        flash(f"فشل السكريبت:{str(e)}", "danger")
        return redirect(url_for("home"))
   ## if scrape_files():
   ##     return "Scrapping completed successfully!"
   ## else:
   ##     return "Scrapping failed"    


@app.route('/update-files')
def update_files():
    """" 
    Trigger web scraping to update documents
    returns: Redirect too home page with flash message
    """
    try:
        if scrape_files():
            flash("Files updated successfully with new downloads!", "success")
        else:
            flash("No new files were downloaded", "info")    
    except Exception as e:
        flash(f"Scraping filed: {str(e)}", "danger")    
    return redirect(url_for("home"))


@app.route('/upload', methods=['POST'])
def upload_file():
    if 'file' not in request.files:
        flash('No file selscted', 'danger')
        return redirect(url_for('home'))
    
    file = request.files['file']
    if file.filename == '':
        flash('No file selected', 'danger')
        return redirect(url_for('home'))
    
    if file and allowed_file(file.filename):
        filename = secure_filename(file.filename)
        try:
            file.save(os.path.join(app.config['UPLOAD_FOLDER'], filename))
            flash('تم رفع الملف بنجاح!', 'success')
        except Exception as e:
            flash(f'Upload failed:{str(e)}', 'danger')
    else:
        flash('Allowed file types are PDF and DOCX', 'danger')

    return redirect(url_for('home'))        


@app.route("/download", methods=["POST"])
def download_button():
        return redirect(url_for('download', filename='sample.pdf'))


@app.route('/download/<filename>')
def download(filename):
    return send_from_directory(DOWNLOAD_FOLDER, filename, as_attachment=True)


@app.route('/classify')
def classify_documents():
    files = os.listdir(DOWNLOAD_FOLDER)
    classifications = {}
    for file in files:
        filepath = os.path.join(DOWNLOAD_FOLDER, file)
        category = classify_file(filepath)
        classifications.setdefault(category, []).append(file)


    return render_template('classify.html', classifications=classifications)

@app.route('/search')
def search():
    query = request.args.get('q','').strip().lower()
    print(f"بحث عن: {query}")
    if not query:
        flash("Please enter a search term", "warning")
        return redirect(url_for('home'))
    
    results = []
   ## queries = [q.strip() for q in query.split() if q.strip()] #تنظيف الكلمات 

    for filename in os.listdir(DOWNLOAD_FOLDER):
        filepath = os.path.join(DOWNLOAD_FOLDER,filename)
        print(f"فحص الملف:{filename}")
        try:
            content = ""  #قراءة المحتوى مع معالجة الاخطاء
            if filename.lower().endswith('.pdf'):
                try:
                    reader = PdfReader(filepath)
                    content = " ".join(page.extract_text() or "" for page in reader.pages).lower()
                except Exception as e:
                    print(f"Error reading PDF {filename}: {str(e)}")
                    continue

            elif filename.lower().endswith('.docx'):
                try:
                    doc = docx.Document(filepath)
                    content = " ".join(para.text for para in doc.paragraphs).lower()
                except Exception as e:
                    print(f"Error reading DOCX {filename}: {str(e)}")
                    continue
            else:
                continue #لتخطي الملفات غير المدعومة

            if all(re.search(r'\b{}\b'.format(re.escape(q)),content) for q in queries): #لبحث اكتر دقة 
                highlighted = content  #تمييز الكلمات كلها
                for q in queries:
                    highlighted = re.sub(r'({})'.format(re.escape(q)),r'<mark style="background-color:yellow">\1</mark>',highlighted, flags=re.IGNORECASE)
                results.append({  # لتمييز النص 
                    'name':filename,
                    'title': extract_title_from_metadata(filepath) or filename,
                    'content': highlighted,
                    'category': classify_file(filepath)
                })

        except Exception as e:
            print(f"General Error with {filename}:{str(e)}")
            continue

    for result in results:
        filepath = os.path.join(DOWNLOAD_FOLDER, result['name'])
        if result['name'].lower().endswith('.pdf'):
            highlighted_path = highlight_text_in_pdf(filepath, query)
            result['highlighted_path'] = highlighted_path
        elif result['name'].lower().endswith('.docx'):
            highlighted_path = highlight_text_in_docx(filepath, query)
            result['highlighted_path'] = highlighted_path

    return render_template('search_results.html', results=results, query=query, files_count=len(results))            

if __name__ == '__main__':
    app.run(debug=True, host='0.0.0.0', port=5000)

